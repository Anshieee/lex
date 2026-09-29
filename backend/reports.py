# backend/reports.py
"""
Monthly Statistics Report Generator for LEX Workbench.
Implements GET /api/reports/monthly per contract §1.2.
"""
import os
import json
import fcntl
from datetime import datetime, timezone
from typing import Optional, List, Dict, Any
from collections import defaultdict

from fastapi import APIRouter, HTTPException, Depends, Query
from fastapi.responses import FileResponse
from pydantic import BaseModel

from backend.auth import require_admin
from backend.tools.audit_logger import get_log_path

router = APIRouter(prefix="/api/reports", tags=["reports"])


class MonthlyReportResponse(BaseModel):
    month: str
    generated_at: str
    source_log: str
    entries_in_month: int
    totals: Dict[str, Any]
    by_model: List[Dict[str, Any]]
    by_task_type: List[Dict[str, Any]]
    daily: List[Dict[str, Any]]
    egress: Dict[str, Any]


def parse_month(month_str: str) -> tuple[int, int]:
    """Parse YYYY-MM string into (year, month). Raises ValueError if invalid."""
    try:
        year, month = map(int, month_str.split("-"))
        if not (1 <= month <= 12):
            raise ValueError("Invalid month")
        return year, month
    except (ValueError, AttributeError):
        raise ValueError("Invalid month format, expected YYYY-MM")


def get_month_range(year: int, month: int) -> tuple[datetime, datetime]:
    """Get start and end datetime for a month (UTC)."""
    start = datetime(year, month, 1, tzinfo=timezone.utc)
    if month == 12:
        end = datetime(year + 1, 1, 1, tzinfo=timezone.utc)
    else:
        end = datetime(year, month + 1, 1, tzinfo=timezone.utc)
    return start, end


def safe_divide(num: float, denom: float) -> Optional[float]:
    """Return num/denom as fraction 0-1, or None if denom is 0."""
    if denom == 0:
        return None
    return round(num / denom, 4)


def load_audit_entries_in_month(year: int, month: int) -> List[Dict[str, Any]]:
    """Load all audit log entries for a given month."""
    start, end = get_month_range(year, month)
    entries = []

    log_path = get_log_path()
    if not os.path.exists(log_path):
        return entries

    with open(log_path, "r") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                entry = json.loads(line)
                ts_str = entry.get("timestamp", "")
                if not ts_str:
                    continue
                ts = datetime.fromisoformat(ts_str.replace("Z", "+00:00"))
                if start <= ts < end:
                    entries.append(entry)
            except (json.JSONDecodeError, KeyError, ValueError):
                continue

    return entries


def generate_monthly_report(entries: List[Dict[str, Any]], year: int, month: int) -> Dict[str, Any]:
    """Generate the monthly report from audit entries."""
    # Totals
    total_tasks = len([e for e in entries if e.get("step_type") == "task_submitted"])
    completed = len([e for e in entries if e.get("step_type") == "task_submitted" and e.get("status") == "success"])
    failed = len([e for e in entries if e.get("step_type") == "task_submitted" and e.get("status") != "success"])
    approvals = len([e for e in entries if e.get("step_type") == "approval" and e.get("status") == "success"])
    rejections = len([e for e in entries if e.get("step_type") == "approval" and e.get("status") == "failed"])

    # By model
    by_model = defaultdict(lambda: {
        "requests": 0, "successes": 0, "total_latency": 0.0,
        "input_tokens": 0, "output_tokens": 0
    })

    for e in entries:
        model = e.get("model_used", "unknown")
        if model and model != "unknown":
            by_model[model]["requests"] += 1
            if e.get("status") == "success":
                by_model[model]["successes"] += 1
            by_model[model]["total_latency"] += e.get("duration_ms", 0)
            # Note: audit log doesn't record token counts, so these remain 0 (will be null in output)

    by_model_list = []
    for model, stats in by_model.items():
        by_model_list.append({
            "model": model,
            "requests": stats["requests"],
            "success_rate": safe_divide(stats["successes"], stats["requests"]),
            "avg_latency_ms": round(stats["total_latency"] / stats["requests"], 1) if stats["requests"] > 0 else None,
            "input_tokens": None,  # Not recorded in audit log
            "output_tokens": None,  # Not recorded in audit log
        })

    # By task type
    by_task_type = defaultdict(lambda: {"count": 0, "successes": 0})

    for e in entries:
        task_type = e.get("step_type", "unknown")
        by_task_type[task_type]["count"] += 1
        if e.get("status") == "success":
            by_task_type[task_type]["successes"] += 1

    by_task_type_list = []
    for task_type, stats in by_task_type.items():
        by_task_type_list.append({
            "task_type": task_type,
            "count": stats["count"],
            "success_rate": safe_divide(stats["successes"], stats["count"]),
        })

    # Daily breakdown
    daily = defaultdict(lambda: {"tasks": 0, "failures": 0})

    for e in entries:
        if e.get("step_type") == "task_submitted":
            ts_str = e.get("timestamp", "")
            if ts_str:
                try:
                    ts = datetime.fromisoformat(ts_str.replace("Z", "+00:00"))
                    date_str = ts.date().isoformat()
                    daily[date_str]["tasks"] += 1
                    if e.get("status") != "success":
                        daily[date_str]["failures"] += 1
                except ValueError:
                    pass

    daily_list = []
    for date_str in sorted(daily.keys()):
        daily_list.append({
            "date": date_str,
            "tasks": daily[date_str]["tasks"],
            "failures": daily[date_str]["failures"],
        })

    # Egress (not recorded in audit log)
    egress = {
        "outbound_connections": None,
        "source": "audit_log (not recorded)"
    }

    return {
        "month": f"{year:04d}-{month:02d}",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_log": get_log_path(),
        "entries_in_month": len(entries),
        "totals": {
            "tasks": total_tasks,
            "completed": completed,
            "failed": failed,
            "success_rate": safe_divide(completed, total_tasks) if total_tasks > 0 else None,
            "approvals": approvals,
            "rejections": rejections,
        },
        "by_model": by_model_list,
        "by_task_type": by_task_type_list,
        "daily": daily_list,
        "egress": egress,
    }


def generate_docx_report(report: Dict[str, Any]) -> str:
    """Generate DOCX report from report dict. Returns path to generated file."""
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Title
    title = doc.add_heading(f"LEX Monthly Statistics Report — {report['month']}", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadata
    doc.add_paragraph(f"Generated at: {report['generated_at']}")
    doc.add_paragraph(f"Source log: {report['source_log']}")
    doc.add_paragraph(f"Entries in month: {report['entries_in_month']}")

    # Summary table
    doc.add_heading("Summary", level=2)
    totals = report["totals"]
    summary_table = doc.add_table(rows=6, cols=2, style="Table Grid")
    summary_data = [
        ("Tasks Submitted", totals["tasks"]),
        ("Completed", totals["completed"]),
        ("Failed", totals["failed"]),
        ("Success Rate", f"{totals['success_rate']:.2%}" if totals["success_rate"] is not None else "N/A"),
        ("Approvals", totals["approvals"]),
        ("Rejections", totals["rejections"]),
    ]
    for i, (label, value) in enumerate(summary_data):
        summary_table.rows[i].cells[0].text = label
        summary_table.rows[i].cells[1].text = str(value)

    # By Model table
    if report["by_model"]:
        doc.add_heading("By Model", level=2)
        model_table = doc.add_table(rows=len(report["by_model"]) + 1, cols=5, style="Table Grid")
        model_table.rows[0].cells[0].text = "Model"
        model_table.rows[0].cells[1].text = "Requests"
        model_table.rows[0].cells[2].text = "Success Rate"
        model_table.rows[0].cells[3].text = "Avg Latency (ms)"
        model_table.rows[0].cells[4].text = "Tokens (In/Out)"
        for i, m in enumerate(report["by_model"]):
            model_table.rows[i + 1].cells[0].text = m["model"]
            model_table.rows[i + 1].cells[1].text = str(m["requests"])
            model_table.rows[i + 1].cells[2].text = f"{m['success_rate']:.2%}" if m["success_rate"] is not None else "N/A"
            model_table.rows[i + 1].cells[3].text = str(m["avg_latency_ms"]) if m["avg_latency_ms"] is not None else "N/A"
            model_table.rows[i + 1].cells[4].text = "Not recorded"

    # By Task Type table
    if report["by_task_type"]:
        doc.add_heading("By Task Type", level=2)
        task_table = doc.add_table(rows=len(report["by_task_type"]) + 1, cols=3, style="Table Grid")
        task_table.rows[0].cells[0].text = "Task Type"
        task_table.rows[0].cells[1].text = "Count"
        task_table.rows[0].cells[2].text = "Success Rate"
        for i, t in enumerate(report["by_task_type"]):
            task_table.rows[i + 1].cells[0].text = t["task_type"]
            task_table.rows[i + 1].cells[1].text = str(t["count"])
            task_table.rows[i + 1].cells[2].text = f"{t['success_rate']:.2%}" if t["success_rate"] is not None else "N/A"

    # Daily table
    if report["daily"]:
        doc.add_heading("Daily Breakdown", level=2)
        daily_table = doc.add_table(rows=len(report["daily"]) + 1, cols=3, style="Table Grid")
        daily_table.rows[0].cells[0].text = "Date"
        daily_table.rows[0].cells[1].text = "Tasks"
        daily_table.rows[0].cells[2].text = "Failures"
        for i, d in enumerate(report["daily"]):
            daily_table.rows[i + 1].cells[0].text = d["date"]
            daily_table.rows[i + 1].cells[1].text = str(d["tasks"])
            daily_table.rows[i + 1].cells[2].text = str(d["failures"])

    # Egress
    doc.add_heading("Egress", level=2)
    egress = report["egress"]
    doc.add_paragraph(f"Outbound connections: {egress['outbound_connections'] if egress['outbound_connections'] is not None else 'Not recorded'}")
    doc.add_paragraph(f"Source: {egress['source']}")

    # Save to file
    filename = f"lex-report-{report['month']}.docx"
    output_path = os.path.join("/tmp", filename)
    doc.save(output_path)
    return output_path


@router.get("/monthly")
async def get_monthly_report(
    month: str = Query(..., description="Month in YYYY-MM format"),
    format: str = Query("json", description="Output format: json or docx"),
    user: dict = Depends(require_admin),
):
    """
    Generate monthly statistics report.
    Admin only. Returns JSON or DOCX.
    """
    # Validate month format
    try:
        year, month_num = parse_month(month)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    # Load entries for the month
    entries = load_audit_entries_in_month(year, month_num)

    # Generate report
    report = generate_monthly_report(entries, year, month_num)

    if format == "json":
        return report
    elif format == "docx":
        docx_path = generate_docx_report(report)
        return FileResponse(
            docx_path,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=f"lex-report-{month}.docx",
        )
    else:
        raise HTTPException(status_code=400, detail="Invalid format, must be 'json' or 'docx'")